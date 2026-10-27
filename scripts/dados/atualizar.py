"""
Atualiza os dados de mercado do site (rodar 1x por mês, depois commitar src/data/mercado/).

    python3 scripts/dados/atualizar.py            # tudo
    python3 scripts/dados/atualizar.py ivgr fipezap

Fontes (ver docs/pesquisa/fontes-de-dados-imoveis.md):
  ivgr      Banco Central, IVG-R (SGS 21340): índice das avaliações dos imóveis financiados.
  fipezap   Fipe/ZAP, série histórica: m² médio de ANÚNCIO, 56 cidades.
  brasil    IBGE, malha das UFs: o desenho do mapa do Brasil.
  itbi-sp   Prefeitura de SP, guias de ITBI pagas (preço REAL), por distrito (GeoSampa).

Os arquivos brutos ficam em .cache/dados/ (fora do git). Precisa de pandas, openpyxl e shapely.
"""

import json
import math
import re
import statistics
import subprocess
import sys
from datetime import date
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[2]
CACHE = RAIZ / ".cache" / "dados"
SAIDA = RAIZ / "src" / "data" / "mercado"
UA = "Mozilla/5.0 (PrimeiraChave dados)"


def baixar(url: str, nome: str, refazer: bool = False) -> Path:
    destino = CACHE / nome
    if destino.exists() and not refazer:
        return destino
    destino.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(["curl", "-sSL", "-m", "900", "-A", UA, "-o", str(destino), url], check=True)
    return destino


def salvar(nome: str, dados) -> None:
    SAIDA.mkdir(parents=True, exist_ok=True)
    (SAIDA / nome).write_text(json.dumps(dados, ensure_ascii=False, separators=(",", ":")) + "\n")
    print(f"  → src/data/mercado/{nome}")


# ——— IVG-R ————————————————————————————————————————————————————————————————


def ivgr():
    url = "https://api.bcb.gov.br/dados/serie/bcdata.sgs.21340/dados?formato=json&dataInicial=01/01/2016"
    serie = json.loads(baixar(url, "ivgr.json", refazer=True).read_text())
    pontos = [[f"{d['data'][6:]}-{d['data'][3:5]}", round(float(d["valor"]), 2)] for d in serie]
    salvar("ivgr.json", {"fonte": "Banco Central do Brasil, IVG-R (SGS 21340)", "base": "mar/2001 = 100", "serie": pontos})


# ——— FipeZap ——————————————————————————————————————————————————————————————

CAPITAIS = {
    # cidade: (UF, lat, lon)
    "São Paulo": ("SP", -23.55, -46.63), "Rio de Janeiro": ("RJ", -22.91, -43.17), "Belo Horizonte": ("MG", -19.92, -43.94),
    "Vitória": ("ES", -20.32, -40.34), "Curitiba": ("PR", -25.43, -49.27), "Florianópolis": ("SC", -27.60, -48.55),
    "Porto Alegre": ("RS", -30.03, -51.23), "Brasília": ("DF", -15.79, -47.88), "Goiânia": ("GO", -16.69, -49.25),
    "Campo Grande": ("MS", -20.46, -54.62), "Cuiabá": ("MT", -15.60, -56.10), "Aracaju": ("SE", -10.91, -37.07),
    "Fortaleza": ("CE", -3.73, -38.52), "João Pessoa": ("PB", -7.12, -34.86), "Maceió": ("AL", -9.67, -35.74),
    "Natal": ("RN", -5.79, -35.21), "Recife": ("PE", -8.05, -34.88), "Salvador": ("BA", -12.97, -38.50),
    "São Luís": ("MA", -2.53, -44.30), "Teresina": ("PI", -5.09, -42.80), "Belém": ("PA", -1.46, -48.50),
    "Manaus": ("AM", -3.12, -60.02),
}


def fipezap():
    import openpyxl

    arq = baixar("https://downloads.fipe.org.br/indices/fipezap/fipezap-serieshistoricas.xlsx", "fipezap.xlsx", refazer=True)
    wb = openpyxl.load_workbook(arq, read_only=True, data_only=True)
    resumo = list(wb["Resumo"].iter_rows(values_only=True))
    ref = resumo[5][3].strftime("%Y-%m")
    cidades = []
    for r in resumo[6:]:
        nome, uf = r[1], r[2]
        if not nome or not uf or nome not in wb.sheetnames:
            continue
        linhas = [x for x in wb[nome].iter_rows(values_only=True) if hasattr(x[1], "year")]
        ult = linhas[-1]
        # colunas da aba da cidade: 17..21 = preço médio venda (Total, 1D, 2D, 3D, 4D); 12 = var. 12 meses total
        por_quartos = {str(q): round(v) for q, v in zip(range(1, 5), ult[18:22]) if isinstance(v, (int, float))}
        cidade = {"nome": nome, "uf": uf, "m2": round(r[6]), "var12": round(r[5] * 100, 1), "quartos": por_quartos}
        if nome in CAPITAIS:
            cidade["capital"] = True
            cidade["lat"], cidade["lon"] = CAPITAIS[nome][1:]
        cidades.append(cidade)
    brasil = next(r for r in resumo[6:] if r[1] == "Índice FipeZAP")
    salvar(
        "fipezap.json",
        {
            "fonte": "Índice FipeZAP (Fipe e Grupo OLX), preço médio de anúncio, média móvel de 3 meses",
            "ref": ref,
            "brasil": {"m2": round(brasil[6]), "var12": round(brasil[5] * 100, 1)},
            "cidades": cidades,
        },
    )


# ——— Desenho dos mapas (SVG) ——————————————————————————————————————————————


def projetor(minx, miny, maxx, maxy, largura):
    """Equiretangular com correção de latitude: bom o bastante para um país ou uma cidade."""
    k = math.cos(math.radians((miny + maxy) / 2))
    escala = largura / ((maxx - minx) * k)
    altura = (maxy - miny) * escala

    def p(lon, lat):
        return round((lon - minx) * k * escala, 1), round((maxy - lat) * escala, 1)

    return p, round(altura)


def caminho(geom, p) -> str:
    partes = []
    polys = [geom] if geom.geom_type == "Polygon" else list(geom.geoms)
    for poly in polys:
        for anel in [poly.exterior, *poly.interiors]:
            pts = [p(x, y) for x, y in anel.coords]
            partes.append("M" + "L".join(f"{x} {y}" for x, y in pts) + "Z")
    return "".join(partes)


def brasil():
    from shapely.geometry import shape

    url = "https://servicodados.ibge.gov.br/api/v3/malhas/paises/BR?formato=application/vnd.geo+json&qualidade=minima&intrarregiao=UF"
    geo = json.loads(baixar(url, "brasil-uf.json").read_text())
    nomes = {x["id"]: x["sigla"] for x in json.loads(baixar("https://servicodados.ibge.gov.br/api/v1/localidades/estados", "ufs.json").read_text())}
    feats = [(nomes[int(f["properties"]["codarea"])], shape(f["geometry"]).simplify(0.06)) for f in geo["features"]]
    minx = min(g.bounds[0] for _, g in feats); miny = min(g.bounds[1] for _, g in feats)
    maxx = max(g.bounds[2] for _, g in feats); maxy = max(g.bounds[3] for _, g in feats)
    p, altura = projetor(minx, miny, maxx, maxy, 600)
    capitais = {nome: p(lon, lat) for nome, (_, lat, lon) in CAPITAIS.items()}
    salvar(
        "mapa-brasil.json",
        {"fonte": "IBGE, malha das UFs", "largura": 600, "altura": altura, "ufs": [{"uf": uf, "d": caminho(g, p)} for uf, g in sorted(feats)], "capitais": capitais},
    )


# ——— ITBI de São Paulo ————————————————————————————————————————————————————

ITBI_PAGINA = "https://prefeitura.sp.gov.br/web/fazenda/w/acesso_a_informacao/31501"
GEOSAMPA = "https://wfs.geosampa.prefeitura.sp.gov.br/geoserver/ows?service=WFS&version=2.0.0&request=GetFeature&outputFormat=application/json&srsName=EPSG:4326&typeNames=geoportal:"
MESES = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]
MIN_DISTRITO = 30  # abaixo disso a mediana não diz muito


def itbi_sp():
    import pandas as pd
    from shapely import STRtree
    from shapely.geometry import shape

    # 1. As duas planilhas mais recentes (ano corrente + anterior) listadas na página da Fazenda.
    html = subprocess.run(["curl", "-sSL", "-A", UA, ITBI_PAGINA], capture_output=True, text=True, check=True).stdout
    links = list(dict.fromkeys(re.findall(r'href="([^"]+GUIAS[^"]+\.xlsx)"', html)))[:2]
    abas = []
    for i, link in enumerate(links):
        xl = pd.ExcelFile(baixar(link.replace(" ", "%20"), f"itbi-sp-{i}.xlsx", refazer=True))
        for aba in xl.sheet_names:
            if re.fullmatch(r"[A-Z]{3}-\d{4}", aba):
                abas.append((int(aba[4:]) * 100 + MESES.index(aba[:3]) + 1, xl, aba))
    abas.sort(key=lambda a: a[0])
    ultimos = abas[-12:]  # 12 meses
    df = pd.concat([xl.parse(aba) for _, xl, aba in ultimos], ignore_index=True)
    ym = lambda n: f"{n // 100}-{n % 100:02d}"
    periodo = {"de": ym(ultimos[0][0]), "ate": ym(ultimos[-1][0])}

    # 2. Só compra e venda residencial do imóvel inteiro, com área e valor plausíveis.
    uso = df["Descrição do uso (IPTU)"].fillna("")
    r = df[(df["Natureza de Transação"] == "1.Compra e venda") & (df["Proporção Transmitida (%)"] == 100) & (uso.str.startswith("APARTAMENTO") | (uso == "RESIDÊNCIA"))].copy()
    r["apto"] = r["Descrição do uso (IPTU)"].str.startswith("APARTAMENTO")
    r["area"] = pd.to_numeric(r["Área Construída (m2)"], errors="coerce")
    r["valor"] = pd.to_numeric(r["Valor de Transação (declarado pelo contribuinte)"], errors="coerce")
    r = r[(r.area >= 20) & (r.area <= 600) & (r.valor >= 50000)]
    r["m2"] = r.valor / r.area
    r = r[(r.m2 >= 1500) & (r.m2 <= 60000)]
    fin = r["Tipo de Financiamento"].astype(str)
    r["financiado"] = fin.str.match(r"^(1|2|3|99)\.")
    r["mcmv"] = fin.str.startswith("2.")

    # 3. Cada compra vai para o distrito pela quadra fiscal (o SQL começa com setor + quadra).
    distritos = json.loads(baixar(GEOSAMPA + "distrito_municipal", "sp-distritos.json").read_text())["features"]
    polys = [shape(f["geometry"]) for f in distritos]
    arvore = STRtree(polys)
    quadra_distrito = {}
    for ini in range(0, 80000, 20000):
        q = json.loads(baixar(GEOSAMPA + f"quadra_fiscal&count=20000&startIndex={ini}&sortBy=cd_identificador", f"sp-quadras-{ini}.json").read_text())
        for f in q["features"]:
            pr, pt = f["properties"], shape(f["geometry"]).representative_point()
            for i in arvore.query(pt):
                if polys[i].contains(pt):
                    quadra_distrito[pr["cd_setor_fiscal"].zfill(3) + pr["cd_quadra_fiscal"].zfill(3)] = i
                    break
    sql = r["N° do Cadastro (SQL)"].astype("Int64").astype(str).str.zfill(11)
    r["d"] = [quadra_distrito.get(s[:6]) for s in sql]
    r = r[r.d.notna()]

    # 4. Desenho e números por distrito.
    from shapely.ops import unary_union

    cidade = unary_union(polys)
    p, altura = projetor(*cidade.bounds, 520)
    saida = []
    for i, f in enumerate(distritos):
        pr = f["properties"]
        x = r[r.d == i]
        item = {"nome": pr["nm_distrito_municipal"].title(), "regiao": pr["nm_regiao_05"], "d": caminho(polys[i].simplify(0.0012), p), "n": int(len(x))}
        if len(x) >= MIN_DISTRITO:
            item.update(
                m2=round(x.m2.median() / 10) * 10,
                valor=round(x.valor.median() / 1000) * 1000,
                financiado=round(x.financiado.mean() * 100),
            )
        saida.append(item)
    saida.sort(key=lambda d: d["nome"])
    salvar(
        "itbi-sp.json",
        {
            "fonte": "Prefeitura de São Paulo, guias de ITBI pagas; distritos do GeoSampa",
            "periodo": periodo,
            "largura": 520,
            "altura": altura,
            "cidade": {
                "n": int(len(r)),
                "m2": round(r.m2.median() / 10) * 10,
                "valor": round(r.valor.median() / 1000) * 1000,
                "financiado": round(r.financiado.mean() * 100),
                "mcmv": round(r.mcmv.mean() * 100, 1),
                "valorFinanciado": round(pd.to_numeric(r[r.financiado]["Valor Financiado"], errors="coerce").median() / 1000) * 1000,
            },
            "distritos": saida,
        },
    )


TAREFAS = {"ivgr": ivgr, "fipezap": fipezap, "brasil": brasil, "itbi-sp": itbi_sp}

if __name__ == "__main__":
    for nome in sys.argv[1:] or TAREFAS:
        print(f"{nome}…")
        TAREFAS[nome]()
    print(f"ok ({date.today():%d/%m/%Y})")
